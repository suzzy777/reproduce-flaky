#!/usr/bin/env python3
"""Automatically retry excluded Java/Maven subjects with compilation repairs.

The pipeline deduplicates test rows by repository URL + commit SHA + Maven
module, then tries the supplied Maven skip flags under Java 8, 11, and 17.
Depending on the existing compilation-error category, it can also rewrite HTTP
URLs in POM files and replace ``-SNAPSHOT`` versions before retrying.

Only Python's standard library is required for CSV input. XLSX category files
are supported when ``openpyxl`` is installed.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import functools
import hashlib
import json
import os
import platform
import re
import shlex
import shutil
import signal
import subprocess
import sys
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence


SKIP_PROPERTIES = [
    "dependency-check.skip=true",
    "gpg.skip=true",
    "failIfNoTests=false",
    "skip.installnodenpm=true",
    "skip.npm=true",
    "skip.yarn=true",
    "license.skip=true",
    "checkstyle.skip=true",
    "rat.skip=true",
    "enforcer.skip=true",
    "animal.sniffer.skip=true",
    "maven.javadoc.skip=true",
    "findbugs.skip=true",
    "spotbugs.skip=true",
    "warbucks.skip=true",
    "modernizer.skip=true",
    "impsort.skip=true",
    "mdep.analyze.skip=true",
    "pgpverify.skip=true",
    "xml.skip=true",
    "cobertura.skip=true",
]

ALIASES = {
    "repo": [
        "project url", "project_url", "repository url", "repository_url",
        "repo url", "repo_url", "github url", "github_url", "git url",
        "git_url", "repository", "repo", "url", "project",
    ],
    "sha": [
        "commit sha", "commit_sha", "git sha", "git_sha", "sha",
        "commit hash", "commit_hash", "commit",
    ],
    "module": [
        "module path", "module_path", "relative module path",
        "relative_module_path", "maven module", "maven_module", "module",
        "relative path", "relative_path",
    ],
    "test": [
        "fully qualified test name", "fully_qualified_test_name", "test name",
        "test_name", "test method", "test_method", "test", "test case",
        "test_case",
    ],
    "category": [
        "compilation error subcategory", "compilation_error_subcategory",
        "error subcategory", "error_subcategory", "subcategory",
        "compilation error category", "compilation_error_category",
        "error category", "error_category", "failure category",
        "failure_category", "category", "error type", "error_type",
    ],
    "fix": [
        "recommended fix", "recommended_fix", "compilation fix",
        "compilation_fix", "repair strategy", "repair_strategy", "strategy",
        "fix",
    ],
    "source": ["dataset", "source", "data source", "data_source"],
}

HTTP_HINTS = (
    "http", "https", "blocked repository", "blocked repo", "maven http",
    "repository mirror", "non-existent http", "nonexistent http",
)
SNAPSHOT_HINTS = (
    "snapshot", "missing dependency", "missing dependencies",
    "dependency resolution", "non-resolvable", "nonresolvable",
    "missing artifact", "required artifact", "import pom", "parent pom",
)


@dataclass
class InputRecord:
    values: dict[str, str]
    input_file: str
    row_number: int
    repo: str
    sha: str
    module: str
    test: str
    category: str
    fix: str
    source: str
    subject_key: tuple[str, str, str]


@dataclass
class Subject:
    subject_id: str
    repo: str
    sha: str
    module: str
    records: list[InputRecord] = field(default_factory=list)
    categories: set[str] = field(default_factory=set)
    fixes: set[str] = field(default_factory=set)
    sources: set[str] = field(default_factory=set)

    @property
    def category_text(self) -> str:
        return " | ".join(sorted(x for x in self.categories if x)) or "Unknown"

    @property
    def fix_text(self) -> str:
        return " | ".join(sorted(x for x in self.fixes if x))


def eprint(*args: object) -> None:
    print(*args, file=sys.stderr, flush=True)


def normalized_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).strip().lower()).strip()


def pick_column(
    headers: Sequence[str], kind: str, explicit: str | None, required: bool = False
) -> str | None:
    if explicit:
        for header in headers:
            if header == explicit or normalized_header(header) == normalized_header(explicit):
                return header
        raise ValueError(f"Column {explicit!r} was not found. Available columns: {list(headers)}")

    normalized = {normalized_header(h): h for h in headers}
    for alias in ALIASES[kind]:
        found = normalized.get(normalized_header(alias))
        if found:
            return found
    if required:
        raise ValueError(
            f"Could not infer the {kind!r} column. Available columns: {list(headers)}. "
            f"Pass --{kind}-column explicitly."
        )
    return None


def read_table(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise RuntimeError(
                f"{path} is an Excel file. Install openpyxl or export it as CSV: "
                "python3 -m pip install openpyxl"
            ) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            return [], []
        headers = [str(v).strip() if v is not None else "" for v in raw_headers]
        records = []
        for row in rows:
            values = list(row) + [None] * max(0, len(headers) - len(row))
            if not any(v not in (None, "") for v in values):
                continue
            records.append(
                {h: "" if v is None else str(v).strip() for h, v in zip(headers, values)}
            )
        return headers, records

    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            return [], []
        headers = [str(h) for h in reader.fieldnames]
        return headers, [
            {str(k): "" if v is None else str(v).strip() for k, v in row.items()}
            for row in reader
        ]


def normalize_repo(value: str) -> str:
    value = value.strip()
    if not value:
        return value
    if re.fullmatch(r"[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+(?:\.git)?", value):
        value = f"https://github.com/{value}"
    elif value.startswith("github.com/"):
        value = f"https://{value}"
    return value.rstrip("/")


def repo_key(value: str) -> str:
    value = normalize_repo(value).lower()
    return value[:-4] if value.endswith(".git") else value


def normalize_module(value: str) -> str:
    value = value.strip().replace("\\", "/")
    if not value or value.lower() in {"root", "project root", "n/a", "none"}:
        return "."
    while value.startswith("./"):
        value = value[2:]
    return value.rstrip("/") or "."


def normalize_test(value: str) -> str:
    return re.sub(r"\s+", "", value.strip().replace("#", "."))


def make_subject_key(repo: str, sha: str, module: str) -> tuple[str, str, str]:
    return repo_key(repo), sha.strip().lower(), normalize_module(module).lower()


def load_input_records(paths: Sequence[Path], args: argparse.Namespace) -> list[InputRecord]:
    output: list[InputRecord] = []
    for path in paths:
        headers, rows = read_table(path)
        if not headers:
            eprint(f"WARNING: {path} is empty; skipping it")
            continue
        repo_col = pick_column(headers, "repo", args.repo_column, required=True)
        sha_col = pick_column(headers, "sha", args.sha_column, required=True)
        module_col = pick_column(headers, "module", args.module_column)
        test_col = pick_column(headers, "test", args.test_column)
        category_col = pick_column(headers, "category", args.category_column)
        fix_col = pick_column(headers, "fix", args.fix_column)
        source_col = pick_column(headers, "source", args.source_column)

        assert repo_col is not None and sha_col is not None
        for row_number, row in enumerate(rows, start=2):
            repo = normalize_repo(row.get(repo_col, ""))
            sha = row.get(sha_col, "").strip()
            if not repo or not sha:
                eprint(f"WARNING: skipping {path.name}:{row_number}; repository or SHA is blank")
                continue
            module = normalize_module(row.get(module_col, "") if module_col else "")
            test = row.get(test_col, "").strip() if test_col else ""
            category = row.get(category_col, "").strip() if category_col else ""
            fix = row.get(fix_col, "").strip() if fix_col else ""
            source = row.get(source_col, "").strip() if source_col else path.stem
            output.append(
                InputRecord(
                    values=row,
                    input_file=str(path),
                    row_number=row_number,
                    repo=repo,
                    sha=sha,
                    module=module,
                    test=test,
                    category=category,
                    fix=fix,
                    source=source,
                    subject_key=make_subject_key(repo, sha, module),
                )
            )
    return output


def load_category_lookup(path: Path | None, args: argparse.Namespace) -> tuple[dict, dict, dict]:
    if path is None:
        return {}, {}, {}
    headers, rows = read_table(path)
    if not headers:
        return {}, {}, {}
    repo_col = pick_column(headers, "repo", args.category_repo_column or args.repo_column, True)
    sha_col = pick_column(headers, "sha", args.category_sha_column or args.sha_column, True)
    module_col = pick_column(headers, "module", args.category_module_column or args.module_column)
    test_col = pick_column(headers, "test", args.category_test_column or args.test_column)
    category_col = pick_column(headers, "category", args.category_category_column or args.category_column, True)
    fix_col = pick_column(headers, "fix", args.category_fix_column or args.fix_column)
    assert repo_col and sha_col and category_col

    by_subject: dict[tuple[str, str, str], list[tuple[str, str]]] = defaultdict(list)
    by_test: dict[tuple[tuple[str, str, str], str], list[tuple[str, str]]] = defaultdict(list)
    test_only_candidates: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for row in rows:
        repo = normalize_repo(row.get(repo_col, ""))
        sha = row.get(sha_col, "").strip()
        if not repo or not sha:
            continue
        module = normalize_module(row.get(module_col, "") if module_col else "")
        category = row.get(category_col, "").strip()
        fix = row.get(fix_col, "").strip() if fix_col else ""
        pair = (category, fix)
        key = make_subject_key(repo, sha, module)
        by_subject[key].append(pair)
        if test_col and row.get(test_col, "").strip():
            test = normalize_test(row[test_col])
            by_test[(key, test)].append(pair)
            test_only_candidates[test].add(pair)

    test_only = {
        test: next(iter(pairs)) for test, pairs in test_only_candidates.items() if len(pairs) == 1
    }
    return dict(by_subject), dict(by_test), test_only


def attach_categories(
    records: list[InputRecord], lookup: tuple[dict, dict, dict]
) -> None:
    by_subject, by_test, test_only = lookup
    for record in records:
        if record.category and record.fix:
            continue
        candidates: list[tuple[str, str]] = []
        if record.test:
            candidates = by_test.get((record.subject_key, normalize_test(record.test)), [])
        if not candidates:
            candidates = by_subject.get(record.subject_key, [])
        if not candidates and record.test:
            pair = test_only.get(normalize_test(record.test))
            candidates = [pair] if pair else []
        categories = sorted({c for c, _ in candidates if c})
        fixes = sorted({f for _, f in candidates if f})
        if categories and not record.category:
            record.category = " | ".join(categories)
        if fixes and not record.fix:
            record.fix = " | ".join(fixes)


def group_subjects(records: list[InputRecord]) -> list[Subject]:
    grouped: dict[tuple[str, str, str], Subject] = {}
    for record in records:
        subject = grouped.get(record.subject_key)
        if subject is None:
            digest = hashlib.sha256("\0".join(record.subject_key).encode()).hexdigest()[:12]
            subject = Subject(digest, record.repo, record.sha, record.module)
            grouped[record.subject_key] = subject
        subject.records.append(record)
        if record.category:
            subject.categories.add(record.category)
        if record.fix:
            subject.fixes.add(record.fix)
        if record.source:
            subject.sources.add(record.source)
    return sorted(grouped.values(), key=lambda s: (repo_key(s.repo), s.sha, s.module))


def strategies_for(subject: Subject, mode: str) -> list[str]:
    strategies = ["flags"]
    if mode == "all":
        return strategies + ["https", "snapshot", "https+snapshot"]
    text = f"{subject.category_text} {subject.fix_text}".lower()
    dependency_general = "dependency" in text or "artifact" in text or "pom" in text
    wants_https = dependency_general or any(hint in text for hint in HTTP_HINTS)
    wants_snapshot = dependency_general or any(hint in text for hint in SNAPSHOT_HINTS)
    if wants_https:
        strategies.append("https")
    if wants_snapshot:
        strategies.append("snapshot")
    if wants_https and wants_snapshot:
        strategies.append("https+snapshot")
    return strategies


def discover_java_home(version: int, explicit: str | None) -> Path | None:
    candidates: list[Path] = []
    if explicit:
        candidates.append(Path(explicit).expanduser())
    for name in (f"JAVA{version}_HOME", f"JAVA_{version}_HOME", f"JDK{version}_HOME"):
        if os.environ.get(name):
            candidates.append(Path(os.environ[name]).expanduser())

    if platform.system() == "Darwin" and Path("/usr/libexec/java_home").exists():
        requested = "1.8" if version == 8 else str(version)
        result = subprocess.run(
            ["/usr/libexec/java_home", "-v", requested], text=True,
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, check=False,
        )
        if result.returncode == 0 and result.stdout.strip():
            candidates.append(Path(result.stdout.strip()))

    linux_names = {
        8: ["java-8-openjdk-amd64", "java-1.8.0-openjdk-amd64", "java-8-openjdk"],
        11: ["java-11-openjdk-amd64", "java-11-openjdk"],
        17: ["java-17-openjdk-amd64", "java-17-openjdk"],
    }
    for base in (Path("/usr/lib/jvm"), Path("/opt/java"), Path("/opt/jdk")):
        for name in linux_names.get(version, []):
            candidates.append(base / name)
        if base.exists():
            candidates.extend(sorted(base.glob(f"*{version}*")))

    seen = set()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved in seen:
            continue
        seen.add(resolved)
        if (resolved / "bin" / "java").is_file():
            return resolved
    return None


def run_checked(command: Sequence[str], cwd: Path | None = None) -> subprocess.CompletedProcess:
    return subprocess.run(
        list(command), cwd=cwd, text=True, stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT, check=False,
    )


def safe_remove_tree(path: Path, allowed_parent: Path) -> None:
    path = path.resolve()
    allowed_parent = allowed_parent.resolve()
    if path == allowed_parent or allowed_parent not in path.parents:
        raise RuntimeError(f"Refusing to remove unsafe path: {path}")
    if path.exists():
        shutil.rmtree(path)


def prepare_checkout(subject: Subject, cache_dir: Path, work_root: Path, fetch: bool) -> Path:
    repo_hash = hashlib.sha256(repo_key(subject.repo).encode()).hexdigest()[:16]
    mirror = cache_dir / "repos" / f"{repo_hash}.git"
    mirror.parent.mkdir(parents=True, exist_ok=True)
    if not mirror.exists():
        eprint(f"  Cloning mirror: {subject.repo}")
        result = run_checked(["git", "clone", "--mirror", subject.repo, str(mirror)])
        if result.returncode:
            raise RuntimeError(f"git clone failed:\n{result.stdout[-4000:]}")
    elif fetch:
        result = run_checked(["git", "--git-dir", str(mirror), "fetch", "--prune", "origin"])
        if result.returncode:
            eprint(f"WARNING: fetch failed for {subject.repo}: {result.stdout[-1000:]}")

    work = work_root / subject.subject_id
    safe_remove_tree(work, work_root)
    work_root.mkdir(parents=True, exist_ok=True)
    result = run_checked(["git", "clone", "--shared", "--no-checkout", str(mirror), str(work)])
    if result.returncode:
        raise RuntimeError(f"git work clone failed:\n{result.stdout[-4000:]}")
    result = run_checked(["git", "checkout", "--detach", subject.sha], cwd=work)
    if result.returncode:
        run_checked(["git", "--git-dir", str(mirror), "fetch", "origin", subject.sha])
        result = run_checked(["git", "checkout", "--detach", subject.sha], cwd=work)
    if result.returncode:
        raise RuntimeError(f"cannot check out {subject.sha}:\n{result.stdout[-4000:]}")
    return work


def reset_checkout(work: Path, sha: str) -> None:
    result = run_checked(["git", "reset", "--hard", sha], cwd=work)
    if result.returncode:
        raise RuntimeError(result.stdout)
    result = run_checked(["git", "clean", "-fdx"], cwd=work)
    if result.returncode:
        raise RuntimeError(result.stdout)


def pom_files(work: Path) -> list[Path]:
    return sorted(
        p for p in work.rglob("pom.xml")
        if ".git" not in p.parts and "target" not in p.parts
    )


def patch_https(work: Path) -> list[dict[str, str]]:
    # Restrict the rewrite to Maven <url> values; schema URLs are untouched.
    pattern = re.compile(r"(<url\b[^>]*>\s*(?:<!\[CDATA\[)?\s*)http://", re.IGNORECASE)
    changes = []
    for pom in pom_files(work):
        text = pom.read_text(encoding="utf-8", errors="surrogateescape")
        updated, count = pattern.subn(r"\1https://", text)
        if count:
            pom.write_text(updated, encoding="utf-8", errors="surrogateescape")
            changes.append({"pom": str(pom.relative_to(work)), "change": f"http_to_https:{count}"})
    return changes


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def direct_child_text(element: ET.Element, name: str) -> str:
    for child in list(element):
        if local_name(child.tag) == name:
            return (child.text or "").strip()
    return ""


def snapshot_coordinates(pom: Path) -> dict[str, set[tuple[str, str]]]:
    """Map a literal SNAPSHOT value to GAVs that use it in one POM."""
    result: dict[str, set[tuple[str, str]]] = defaultdict(set)
    try:
        root = ET.parse(pom).getroot()
    except (ET.ParseError, OSError):
        return result
    properties: dict[str, str] = {}
    for child in list(root):
        if local_name(child.tag) == "properties":
            for prop in list(child):
                properties[local_name(prop.tag)] = (prop.text or "").strip()
    for element in root.iter():
        if local_name(element.tag) not in {"dependency", "plugin", "parent", "extension"}:
            continue
        group = direct_child_text(element, "groupId")
        artifact = direct_child_text(element, "artifactId")
        version = direct_child_text(element, "version")
        if not group and local_name(element.tag) == "plugin":
            group = "org.apache.maven.plugins"
        property_match = re.fullmatch(r"\$\{([^}]+)}", version)
        literal = properties.get(property_match.group(1), "") if property_match else version
        if group and artifact and literal.upper().endswith("-SNAPSHOT"):
            result[literal].add((group, artifact))
    return result


def stable_version_key(version: str) -> tuple:
    parts = re.split(r"[._+-]", version.lower())
    return tuple((0, int(p)) if p.isdigit() else (1, p) for p in parts)


def is_stable_version(version: str) -> bool:
    return not re.search(
        r"(?:snapshot|alpha|beta|milestone|preview|\brc\d*|\bcr\d*|(?:^|[.-])m\d+|[.-]ea(?:[.-]|$))",
        version,
        re.IGNORECASE,
    )


def numeric_version(version: str) -> tuple[int, ...]:
    return tuple(int(x) for x in re.findall(r"\d+", version))


def choose_closest_version(snapshot: str, versions: Iterable[str]) -> str | None:
    base = re.sub(r"-SNAPSHOT$", "", snapshot, flags=re.IGNORECASE)
    stable = sorted({v for v in versions if is_stable_version(v)}, key=stable_version_key)
    if base in stable:
        return base
    if not stable:
        return None
    target = numeric_version(base)
    if not target:
        return stable[-1]
    same_minor = [v for v in stable if numeric_version(v)[:2] == target[:2]]
    same_major = [v for v in stable if numeric_version(v)[:1] == target[:1]]
    candidates = same_minor or same_major or stable
    below = [v for v in candidates if numeric_version(v) <= target]
    return max(below, key=stable_version_key) if below else min(candidates, key=stable_version_key)


@functools.lru_cache(maxsize=4096)
def central_versions(group: str, artifact: str, timeout: int = 20) -> list[str]:
    query = f'g:"{group}" AND a:"{artifact}"'
    params = urllib.parse.urlencode({"q": query, "core": "gav", "rows": "200", "wt": "json"})
    request = urllib.request.Request(
        f"https://search.maven.org/solrsearch/select?{params}",
        headers={"User-Agent": "compile-repair-pipeline/1.0"},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = json.load(response)
    return [doc["v"] for doc in payload.get("response", {}).get("docs", []) if doc.get("v")]


def patch_snapshots(work: Path, resolution: str) -> list[dict[str, str]]:
    poms = pom_files(work)
    # Only change complete XML element values. This covers <version> and
    # version properties without rewriting comments, SCM tags, or arbitrary
    # strings that merely contain the word SNAPSHOT.
    value_pattern = re.compile(
        r"(?P<prefix>>\s*)(?P<value>[A-Za-z0-9_.+\-]+-SNAPSHOT)(?P<suffix>\s*<)",
        re.IGNORECASE,
    )
    tokens: set[str] = set()
    coordinates: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for pom in poms:
        text = pom.read_text(encoding="utf-8", errors="surrogateescape")
        tokens.update(match.group("value") for match in value_pattern.finditer(text))
        for token, gavs in snapshot_coordinates(pom).items():
            coordinates[token].update(gavs)

    replacements: dict[str, str] = {}
    for token in sorted(tokens):
        stripped = re.sub(r"-SNAPSHOT$", "", token, flags=re.IGNORECASE)
        replacement = stripped
        if resolution == "central" and coordinates.get(token):
            central_choices = []
            for group, artifact in sorted(coordinates[token]):
                try:
                    choice = choose_closest_version(token, central_versions(group, artifact))
                except Exception as exc:  # Network errors should not stop the local repair.
                    eprint(f"    Maven Central lookup failed for {group}:{artifact}: {exc}")
                    choice = None
                if choice:
                    central_choices.append(choice)
            # Keep a project-wide token consistent. If GAVs disagree, stripping
            # -SNAPSHOT is the least surprising deterministic fallback.
            if central_choices and len(set(central_choices)) == 1:
                replacement = central_choices[0]
        replacements[token] = replacement

    changes = []
    for pom in poms:
        text = pom.read_text(encoding="utf-8", errors="surrogateescape")
        used_counts: Counter[tuple[str, str]] = Counter()

        def replace_value(match: re.Match[str]) -> str:
            old = match.group("value")
            # Preserve the exact original spelling for values not selected.
            new = next((v for k, v in replacements.items() if k.lower() == old.lower()), old)
            if new != old:
                used_counts[(old, new)] += 1
            return f"{match.group('prefix')}{new}{match.group('suffix')}"

        updated = value_pattern.sub(replace_value, text)
        used = [f"{old}->{new} ({count})" for (old, new), count in used_counts.items()]
        if updated != text:
            pom.write_text(updated, encoding="utf-8", errors="surrogateescape")
            changes.append({"pom": str(pom.relative_to(work)), "change": "; ".join(used)})
    return changes


def apply_strategy(work: Path, strategy: str, snapshot_resolution: str) -> list[dict[str, str]]:
    changes: list[dict[str, str]] = []
    if "https" in strategy:
        changes.extend(patch_https(work))
    if "snapshot" in strategy:
        changes.extend(patch_snapshots(work, snapshot_resolution))
    return changes


def module_arguments(module: str, work: Path) -> list[str]:
    if module == ".":
        if not (work / "pom.xml").is_file():
            raise RuntimeError("root pom.xml was not found")
        return []
    if module.startswith(":"):
        return ["-pl", module, "-am"]
    module_path = (work / module).resolve()
    if work.resolve() not in module_path.parents:
        raise RuntimeError(f"unsafe module path outside repository: {module}")
    if not (module_path / "pom.xml").is_file():
        raise RuntimeError(f"module POM was not found: {module}/pom.xml")
    return ["-pl", module, "-am"]


def maven_command(work: Path, module: str, args: argparse.Namespace) -> list[str]:
    if args.maven:
        executable = shlex.split(args.maven)
    elif (work / "mvnw").is_file():
        executable = ["bash", str(work / "mvnw")]
    else:
        executable = ["mvn"]
    command = executable + ["-B", "clean", "install", "-DskipTests", "-U"]
    command.extend(module_arguments(module, work))
    command.extend(f"-D{prop}" for prop in SKIP_PROPERTIES)
    command.extend(args.extra_maven_arg or [])
    if args.maven_repo:
        command.append(f"-Dmaven.repo.local={Path(args.maven_repo).expanduser().resolve()}")
    return command


def run_build(
    command: Sequence[str], work: Path, java_home: Path, log_path: Path, timeout: int
) -> tuple[str, int | None, float]:
    env = os.environ.copy()
    env["JAVA_HOME"] = str(java_home)
    env["PATH"] = f"{java_home / 'bin'}{os.pathsep}{env.get('PATH', '')}"
    started = time.monotonic()
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("w", encoding="utf-8", errors="replace") as log:
        log.write(f"JAVA_HOME={java_home}\n")
        log.write(f"COMMAND={shlex.join(command)}\n\n")
        log.flush()
        try:
            process = subprocess.Popen(
                list(command), cwd=work, env=env, stdout=log,
                stderr=subprocess.STDOUT, text=True, start_new_session=True,
            )
        except FileNotFoundError as exc:
            log.write(f"\nEXECUTION ERROR: {exc}\n")
            return "MAVEN_NOT_FOUND", None, time.monotonic() - started
        try:
            exit_code = process.wait(timeout=timeout)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid, signal.SIGTERM)
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
                process.wait()
            log.write(f"\nTIMEOUT after {timeout} seconds\n")
            return "TIMEOUT", None, time.monotonic() - started
    return ("SUCCESS" if exit_code == 0 else "BUILD_FAILED"), exit_code, time.monotonic() - started


def log_summary(path: Path, max_chars: int = 1800) -> str:
    try:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return ""
    interesting = [
        line.strip() for line in lines[-500:]
        if re.search(r"\[ERROR]|BUILD FAILURE|COMPILATION ERROR|Could not resolve|Non-resolvable|Fatal error", line, re.I)
    ]
    chosen = interesting[-12:] if interesting else lines[-8:]
    return " || ".join(chosen)[-max_chars:]


def write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_plan(path: Path, subjects: Sequence[Subject], mode: str) -> None:
    rows = []
    for subject in subjects:
        rows.append({
            "subject_id": subject.subject_id,
            "project_url": subject.repo,
            "sha": subject.sha,
            "module": subject.module,
            "sources": " | ".join(sorted(subject.sources)),
            "categories": subject.category_text,
            "recommended_fix": subject.fix_text,
            "test_rows": len(subject.records),
            "strategy_order": " -> ".join(strategies_for(subject, mode)),
        })
    write_csv(
        path,
        ["subject_id", "project_url", "sha", "module", "sources", "categories",
         "recommended_fix", "test_rows", "strategy_order"],
        rows,
    )


def save_patch(work: Path, path: Path) -> None:
    result = run_checked(["git", "diff", "--", "*.xml"], cwd=work)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(result.stdout, encoding="utf-8")


def execute(args: argparse.Namespace) -> int:
    input_paths = [Path(p).expanduser().resolve() for p in args.input]
    category_path = Path(args.category_file).expanduser().resolve() if args.category_file else None
    for path in input_paths + ([category_path] if category_path else []):
        if path and not path.is_file():
            raise FileNotFoundError(path)

    output_dir = Path(args.output).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    records = load_input_records(input_paths, args)
    if not records:
        raise RuntimeError("No usable input rows were found")
    attach_categories(records, load_category_lookup(category_path, args))
    all_subjects = group_subjects(records)
    subjects = all_subjects[: args.limit] if args.limit else all_subjects
    write_plan(output_dir / "plan.csv", all_subjects, args.strategy_mode)

    print(
        f"Loaded {len(records)} test rows and deduplicated them into "
        f"{len(all_subjects)} project/SHA/module subjects."
    )
    if args.limit:
        print(f"Pilot limit: processing the first {len(subjects)} subjects.")
    if args.plan_only:
        print(f"Plan written to {output_dir / 'plan.csv'}")
        return 0

    java_homes = {
        8: discover_java_home(8, args.java8_home),
        11: discover_java_home(11, args.java11_home),
        17: discover_java_home(17, args.java17_home),
    }
    print("Java homes: " + ", ".join(f"{v}={p or 'NOT FOUND'}" for v, p in java_homes.items()))

    attempts: list[dict] = []
    summaries: dict[str, dict] = {}
    work_root = output_dir / "work"
    cache_dir = Path(args.cache_dir).expanduser().resolve() if args.cache_dir else output_dir / "cache"
    available_java = [(v, p) for v, p in java_homes.items() if p]

    for index, subject in enumerate(subjects, start=1):
        print(f"[{index}/{len(subjects)}] {subject.repo} @ {subject.sha} [{subject.module}]")
        summary = {
            "subject_id": subject.subject_id,
            "project_url": subject.repo,
            "sha": subject.sha,
            "module": subject.module,
            "categories": subject.category_text,
            "test_rows": len(subject.records),
            "status": "UNRESOLVED",
            "successful_strategy": "",
            "successful_java": "",
            "successful_log": "",
            "attempt_count": 0,
            "message": "",
        }
        summaries[subject.subject_id] = summary
        if not available_java:
            summary["status"] = "NO_JAVA_FOUND"
            summary["message"] = "Provide --java8-home, --java11-home, and/or --java17-home"
            continue
        work: Path | None = None
        try:
            work = prepare_checkout(subject, cache_dir, work_root, args.fetch)
            if args.init_submodules:
                result = run_checked(["git", "submodule", "update", "--init", "--recursive"], cwd=work)
                if result.returncode:
                    raise RuntimeError(f"submodule initialization failed:\n{result.stdout[-4000:]}")
            module_arguments(subject.module, work)

            success = False
            for strategy in strategies_for(subject, args.strategy_mode):
                reset_checkout(work, subject.sha)
                changes = apply_strategy(work, strategy, args.snapshot_resolution)
                if strategy != "flags" and not changes:
                    attempts.append({
                        "subject_id": subject.subject_id, "project_url": subject.repo,
                        "sha": subject.sha, "module": subject.module,
                        "categories": subject.category_text, "strategy": strategy,
                        "java_version": "", "java_home": "", "command": "",
                        "started_at": dt.datetime.now(dt.timezone.utc).isoformat(),
                        "duration_seconds": "0", "exit_code": "",
                        "status": "NOT_APPLICABLE", "changed_poms": "0",
                        "log_file": "", "error_summary": "No matching POM content to change",
                    })
                    continue

                for java_version, java_home in available_java:
                    assert java_home is not None
                    # Every attempt starts from the exact SHA, then receives only
                    # the repair represented by this strategy.
                    reset_checkout(work, subject.sha)
                    changes = apply_strategy(work, strategy, args.snapshot_resolution)
                    command = maven_command(work, subject.module, args)
                    attempt_no = len([a for a in attempts if a["subject_id"] == subject.subject_id]) + 1
                    log_path = output_dir / "logs" / subject.subject_id / f"{attempt_no:02d}_{strategy}_java{java_version}.log"
                    started_at = dt.datetime.now(dt.timezone.utc).isoformat()
                    print(f"  {strategy}, Java {java_version}")
                    status, exit_code, duration = run_build(
                        command, work, java_home, log_path, args.timeout
                    )
                    row = {
                        "subject_id": subject.subject_id,
                        "project_url": subject.repo,
                        "sha": subject.sha,
                        "module": subject.module,
                        "categories": subject.category_text,
                        "strategy": strategy,
                        "java_version": java_version,
                        "java_home": str(java_home),
                        "command": shlex.join(command),
                        "started_at": started_at,
                        "duration_seconds": f"{duration:.1f}",
                        "exit_code": "" if exit_code is None else exit_code,
                        "status": status,
                        "changed_poms": len({c["pom"] for c in changes}),
                        "log_file": str(log_path),
                        "error_summary": "" if status == "SUCCESS" else log_summary(log_path),
                    }
                    attempts.append(row)
                    summary["attempt_count"] += 1
                    # Checkpoint attempt data after every build.
                    write_csv(output_dir / "attempts.csv", ATTEMPT_FIELDS, attempts)
                    if status == "SUCCESS":
                        patch_path = output_dir / "patches" / f"{subject.subject_id}.patch"
                        save_patch(work, patch_path)
                        summary.update({
                            "status": "REPAIRED",
                            "successful_strategy": strategy,
                            "successful_java": java_version,
                            "successful_log": str(log_path),
                            "message": str(patch_path) if patch_path.stat().st_size else "No POM patch required",
                        })
                        print(f"  SUCCESS with {strategy} on Java {java_version}")
                        success = True
                        break
                if success:
                    break
        except Exception as exc:
            summary["status"] = "PIPELINE_ERROR"
            summary["message"] = str(exc)
            eprint(f"  ERROR: {exc}")
        finally:
            if work and work.exists() and not (args.keep_successful and summary["status"] == "REPAIRED"):
                safe_remove_tree(work, work_root)

    # Include subjects omitted by --limit in the mapped results.
    for subject in all_subjects:
        if subject.subject_id not in summaries:
            summaries[subject.subject_id] = {
                "subject_id": subject.subject_id, "project_url": subject.repo,
                "sha": subject.sha, "module": subject.module,
                "categories": subject.category_text, "test_rows": len(subject.records),
                "status": "NOT_SELECTED", "successful_strategy": "",
                "successful_java": "", "successful_log": "", "attempt_count": 0,
                "message": "Excluded by --limit",
            }

    summary_rows = [summaries[s.subject_id] for s in all_subjects]
    write_csv(output_dir / "subjects.csv", SUBJECT_FIELDS, summary_rows)
    write_csv(output_dir / "attempts.csv", ATTEMPT_FIELDS, attempts)

    original_fields = []
    for record in records:
        for name in record.values:
            if name not in original_fields:
                original_fields.append(name)
    mapped_rows = []
    for record in records:
        digest = hashlib.sha256("\0".join(record.subject_key).encode()).hexdigest()[:12]
        result = summaries[digest]
        row = dict(record.values)
        row.update({
            "pipeline_input_file": record.input_file,
            "pipeline_subject_id": digest,
            "pipeline_category": record.category or "Unknown",
            "compile_status": result["status"],
            "compile_strategy": result["successful_strategy"],
            "compile_java": result["successful_java"],
            "compile_log": result["successful_log"],
            "compile_message": result["message"],
        })
        mapped_rows.append(row)
    mapped_extra = [
        "pipeline_input_file", "pipeline_subject_id", "pipeline_category",
        "compile_status", "compile_strategy", "compile_java", "compile_log",
        "compile_message",
    ]
    write_csv(output_dir / "tests_with_compile_results.csv", original_fields + mapped_extra, mapped_rows)

    repaired = sum(1 for row in summary_rows if row["status"] == "REPAIRED")
    processed = sum(1 for row in summary_rows if row["status"] != "NOT_SELECTED")
    print(f"Finished: {repaired}/{processed} processed subjects repaired.")
    print(f"Results: {output_dir / 'subjects.csv'}")
    return 0


ATTEMPT_FIELDS = [
    "subject_id", "project_url", "sha", "module", "categories", "strategy",
    "java_version", "java_home", "command", "started_at", "duration_seconds",
    "exit_code", "status", "changed_poms", "log_file", "error_summary",
]
SUBJECT_FIELDS = [
    "subject_id", "project_url", "sha", "module", "categories", "test_rows",
    "status", "successful_strategy", "successful_java", "successful_log",
    "attempt_count", "message",
]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Retry excluded Maven subjects using categorized compilation repairs.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""Examples:
  # Inspect column inference, deduplication, and rule routing only
  python3 compile_repair_pipeline.py --input idoft.csv jira.csv --plan-only

  # Manually pilot ten unique project/SHA/module subjects
  python3 compile_repair_pipeline.py --input idoft.csv jira.csv \\
      --category-file categorized_errors.xlsx --limit 10 \\
      --java8-home /path/to/jdk8 --java11-home /path/to/jdk11 \\
      --java17-home /path/to/jdk17

  # Try every POM repair even when a category is unknown
  python3 compile_repair_pipeline.py --input excluded.csv --strategy-mode all
""",
    )
    parser.add_argument("--input", nargs="+", required=True, help="One or more excluded-test CSV/XLSX files")
    parser.add_argument("--category-file", help="Optional CSV/XLSX containing existing error categories")
    parser.add_argument("--output", default="compile_repair_output", help="Output directory")
    parser.add_argument("--cache-dir", help="Persistent bare Git mirror cache (default: OUTPUT/cache)")
    parser.add_argument("--limit", type=int, help="Process only the first N unique subjects for a pilot")
    parser.add_argument("--plan-only", action="store_true", help="Write plan.csv without cloning or building")
    parser.add_argument("--strategy-mode", choices=["category", "all"], default="category")
    parser.add_argument(
        "--snapshot-resolution", choices=["strip", "central"], default="strip",
        help="Remove -SNAPSHOT, or query Maven Central for a nearby stable version",
    )
    parser.add_argument("--timeout", type=int, default=1800, help="Seconds allowed per Maven attempt")
    parser.add_argument("--fetch", action="store_true", help="Refresh existing cached Git mirrors")
    parser.add_argument("--init-submodules", action="store_true", help="Initialize Git submodules")
    parser.add_argument("--keep-successful", action="store_true", help="Keep successful checkout under OUTPUT/work")
    parser.add_argument("--maven", help="Maven command override, e.g. /opt/apache-maven/bin/mvn")
    parser.add_argument("--maven-repo", help="Explicit Maven local repository directory")
    parser.add_argument(
        "--extra-maven-arg", action="append", default=[],
        help="Additional Maven argument; repeat this option for multiple arguments",
    )
    parser.add_argument("--java8-home")
    parser.add_argument("--java11-home")
    parser.add_argument("--java17-home")

    # Input column overrides. Usually unnecessary because common names are inferred.
    for kind in ("repo", "sha", "module", "test", "category", "fix", "source"):
        parser.add_argument(f"--{kind}-column")
    for kind in ("repo", "sha", "module", "test", "category", "fix"):
        parser.add_argument(f"--category-{kind}-column")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if args.limit is not None and args.limit <= 0:
        parser.error("--limit must be greater than zero")
    if args.timeout <= 0:
        parser.error("--timeout must be greater than zero")
    try:
        return execute(args)
    except KeyboardInterrupt:
        eprint("Interrupted")
        return 130
    except Exception as exc:
        eprint(f"ERROR: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
